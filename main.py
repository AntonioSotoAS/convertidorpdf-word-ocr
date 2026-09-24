import os
import re
import shutil
import threading
import tkinter as tk

from tkinter import filedialog, messagebox, ttk

import fitz
import cv2
import numpy as np
import pytesseract

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# CONFIGURACIÓN
# ============================================================

OCR_LANGUAGE = "spa"

# 300 suele ser suficiente.
# 350 mejora algo el OCR pero consume más RAM y demora más.
DPI = 300


# ============================================================
# CONFIGURAR TESSERACT AUTOMÁTICAMENTE
# ============================================================

def configurar_tesseract():

    # Si está agregado al PATH
    tesseract_path = shutil.which("tesseract")

    if tesseract_path:
        pytesseract.pytesseract.tesseract_cmd = tesseract_path
        return tesseract_path

    posibles_rutas = [
        os.path.expandvars(
            r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"
        ),

        os.path.expandvars(
            r"%LOCALAPPDATA%\Tesseract-OCR\tesseract.exe"
        ),

        r"C:\Program Files\Tesseract-OCR\tesseract.exe",

        r"C:\Program Files\PDF24\tesseract\tesseract.exe",

        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]

    for ruta in posibles_rutas:

        if os.path.isfile(ruta):

            pytesseract.pytesseract.tesseract_cmd = ruta

            return ruta

    return None


TESSERACT_PATH = configurar_tesseract()


# ============================================================
# LIMPIEZA DE TEXTO
# ============================================================

def limpiar_texto(texto):

    if not texto:
        return ""

    texto = texto.replace("\x00", "")

    texto = texto.replace("\r\n", "\n")
    texto = texto.replace("\r", "\n")

    # Eliminar espacios sobrantes
    lineas = []

    for linea in texto.split("\n"):

        linea = linea.strip()

        linea = re.sub(
            r"[ \t]+",
            " ",
            linea
        )

        lineas.append(linea)

    texto = "\n".join(lineas)

    # Evitar demasiados saltos vacíos
    texto = re.sub(
        r"\n{3,}",
        "\n\n",
        texto
    )

    return texto.strip()


# ============================================================
# DETECTAR TÍTULOS
# ============================================================

def parece_titulo(texto):

    texto = texto.strip()

    if not texto:
        return False

    if len(texto) > 120:
        return False

    letras = [
        c for c in texto
        if c.isalpha()
    ]

    if not letras:
        return False

    mayusculas = sum(
        1 for c in letras
        if c.isupper()
    )

    porcentaje = mayusculas / len(letras)

    # Mayormente mayúsculas
    if porcentaje >= 0.70:
        return True

    # Algunos encabezados comunes
    patrones = [
        r"^CAP[IÍ]TULO",
        r"^T[IÍ]TULO",
        r"^SECCI[ÓO]N",
        r"^PARTE",
        r"^LIBRO",
        r"^INTRODUCCI[ÓO]N",
        r"^PR[ÓO]LOGO",
        r"^CONCLUSIONES",
        r"^BIBLIOGRAF[IÍ]A",
    ]

    for patron in patrones:

        if re.search(
            patron,
            texto,
            re.IGNORECASE
        ):
            return True

    return False


# ============================================================
# PREPROCESAR IMAGEN
# ============================================================

def preparar_imagen(pixmap):

    datos = np.frombuffer(
        pixmap.samples,
        dtype=np.uint8
    )

    imagen = datos.reshape(
        pixmap.height,
        pixmap.width,
        pixmap.n
    )

    # RGBA
    if pixmap.n == 4:

        imagen = cv2.cvtColor(
            imagen,
            cv2.COLOR_RGBA2RGB
        )

    # Gris
    gris = cv2.cvtColor(
        imagen,
        cv2.COLOR_RGB2GRAY
    )

    # Reducción de ruido
    gris = cv2.GaussianBlur(
        gris,
        (3, 3),
        0
    )

    # Binarización automática
    _, binaria = cv2.threshold(
        gris,
        0,
        255,
        cv2.THRESH_BINARY +
        cv2.THRESH_OTSU
    )

    return binaria


# ============================================================
# OCR
# ============================================================

def ejecutar_ocr(imagen):

    config = (
        "--oem 3 "
        "--psm 6 "
        "-c preserve_interword_spaces=1"
    )

    texto = pytesseract.image_to_string(
        imagen,
        lang=OCR_LANGUAGE,
        config=config
    )

    return limpiar_texto(texto)


# ============================================================
# RECONSTRUCCIÓN DE PÁRRAFOS
# ============================================================

def construir_parrafos(texto):

    lineas = texto.split("\n")

    parrafos = []

    actual = ""

    for linea in lineas:

        linea = linea.strip()

        # Línea vacía
        if not linea:

            if actual:

                parrafos.append(
                    actual.strip()
                )

                actual = ""

            continue

        # Si parece título
        if parece_titulo(linea):

            if actual:

                parrafos.append(
                    actual.strip()
                )

                actual = ""

            parrafos.append(
                ("__TITULO__", linea)
            )

            continue

        # Palabra cortada por guion
        if actual.endswith("-"):

            actual = (
                actual[:-1] +
                linea
            )

            continue

        if not actual:

            actual = linea

        else:

            actual += " " + linea

        # Final de párrafo probable
        if linea.endswith(
            (".", ":", ";", "!", "?")
        ):

            parrafos.append(
                actual.strip()
            )

            actual = ""

    if actual:

        parrafos.append(
            actual.strip()
        )

    return parrafos


# ============================================================
# CONFIGURACIÓN WORD
# ============================================================

def configurar_word(documento):

    seccion = documento.sections[0]

    seccion.top_margin = Cm(2.5)
    seccion.bottom_margin = Cm(2.5)
    seccion.left_margin = Cm(2.5)
    seccion.right_margin = Cm(2.5)

    estilo = documento.styles["Normal"]

    estilo.font.name = (
        "Times New Roman"
    )

    estilo.font.size = Pt(11)


# ============================================================
# AGREGAR TEXTO AL WORD
# ============================================================

def agregar_texto_word(documento, texto):

    elementos = construir_parrafos(
        texto
    )

    for elemento in elementos:

        # Título
        if (
            isinstance(elemento, tuple)
            and elemento[0] == "__TITULO__"
        ):

            titulo = elemento[1]

            parrafo = (
                documento.add_paragraph()
            )

            parrafo.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = parrafo.add_run(
                titulo
            )

            run.bold = True
            run.font.name = (
                "Times New Roman"
            )

            run.font.size = Pt(12)

            continue

        # Texto normal
        parrafo = (
            documento.add_paragraph()
        )

        parrafo.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
        )

        formato = (
            parrafo.paragraph_format
        )

        formato.first_line_indent = (
            Cm(1.25)
        )

        formato.space_after = Pt(5)

        formato.line_spacing = 1.15

        run = parrafo.add_run(
            elemento
        )

        run.font.name = (
            "Times New Roman"
        )

        run.font.size = Pt(11)


# ============================================================
# APLICACIÓN
# ============================================================

class AplicacionOCR:

    def __init__(self, root):

        self.root = root

        self.root.title(
            "PDF escaneado a Word - OCR"
        )

        self.root.geometry(
            "800x520"
        )

        self.root.minsize(
            700,
            480
        )

        self.pdf_path = (
            tk.StringVar()
        )

        self.word_path = (
            tk.StringVar()
        )

        self.estado = tk.StringVar(
            value="Inicializando..."
        )

        self.progreso = (
            tk.DoubleVar(
                value=0
            )
        )

        self.cancelar_proceso = (
            False
        )

        self.procesando = False

        self.crear_interfaz()

        self.root.after(
            300,
            self.verificar_tesseract
        )


    # ========================================================
    # INTERFAZ
    # ========================================================

    def crear_interfaz(self):

        frame = ttk.Frame(
            self.root,
            padding=20
        )

        frame.pack(
            fill="both",
            expand=True
        )

        titulo = ttk.Label(
            frame,
            text=(
                "Convertidor PDF escaneado → Word"
            ),
            font=(
                "Segoe UI",
                18,
                "bold"
            )
        )

        titulo.pack(
            pady=(0, 5)
        )

        descripcion = ttk.Label(
            frame,
            text=(
                "Convierte PDFs escaneados "
                "en documentos Word editables "
                "utilizando OCR en español."
            ),
            font=(
                "Segoe UI",
                10
            )
        )

        descripcion.pack(
            pady=(0, 20)
        )

        # PDF
        frame_pdf = ttk.LabelFrame(
            frame,
            text="PDF de entrada",
            padding=10
        )

        frame_pdf.pack(
            fill="x",
            pady=7
        )

        ttk.Entry(
            frame_pdf,
            textvariable=self.pdf_path
        ).pack(
            side="left",
            fill="x",
            expand=True,
            padx=(0, 10)
        )

        ttk.Button(
            frame_pdf,
            text="Seleccionar PDF",
            command=self.seleccionar_pdf
        ).pack(
            side="right"
        )

        # WORD
        frame_word = ttk.LabelFrame(
            frame,
            text="Word de salida",
            padding=10
        )

        frame_word.pack(
            fill="x",
            pady=7
        )

        ttk.Entry(
            frame_word,
            textvariable=self.word_path
        ).pack(
            side="left",
            fill="x",
            expand=True,
            padx=(0, 10)
        )

        ttk.Button(
            frame_word,
            text="Guardar como...",
            command=self.seleccionar_word
        ).pack(
            side="right"
        )

        # PROGRESO
        frame_progreso = (
            ttk.LabelFrame(
                frame,
                text="Progreso",
                padding=10
            )
        )

        frame_progreso.pack(
            fill="x",
            pady=(20, 10)
        )

        self.barra = (
            ttk.Progressbar(
                frame_progreso,
                variable=self.progreso,
                maximum=100
            )
        )

        self.barra.pack(
            fill="x",
            pady=5
        )

        ttk.Label(
            frame_progreso,
            textvariable=self.estado
        ).pack(
            anchor="w",
            pady=5
        )

        # BOTONES
        frame_botones = ttk.Frame(
            frame
        )

        frame_botones.pack(
            pady=20
        )

        self.btn_convertir = (
            ttk.Button(
                frame_botones,
                text="Iniciar conversión",
                command=self.iniciar_conversion
            )
        )

        self.btn_convertir.pack(
            side="left",
            padx=5
        )

        self.btn_cancelar = (
            ttk.Button(
                frame_botones,
                text="Cancelar",
                command=self.cancelar,
                state="disabled"
            )
        )

        self.btn_cancelar.pack(
            side="left",
            padx=5
        )


    # ========================================================
    # VERIFICAR TESSERACT
    # ========================================================

    def verificar_tesseract(self):

        if not TESSERACT_PATH:

            self.estado.set(
                "No se encontró Tesseract OCR."
            )

            messagebox.showerror(
                "Tesseract no encontrado",
                (
                    "No se encontró Tesseract OCR.\n\n"
                    "Instale Tesseract antes "
                    "de utilizar la aplicación."
                )
            )

            return

        try:

            version = (
                pytesseract
                .get_tesseract_version()
            )

            idiomas = (
                pytesseract
                .get_languages(
                    config=""
                )
            )

            if OCR_LANGUAGE not in idiomas:

                self.estado.set(
                    "Tesseract encontrado, "
                    "pero falta español (spa)."
                )

                messagebox.showwarning(
                    "Idioma español",
                    (
                        "Tesseract fue encontrado en:\n\n"
                        f"{TESSERACT_PATH}\n\n"
                        "Pero no se encontró "
                        "el idioma español 'spa'."
                    )
                )

                return

            self.estado.set(
                (
                    "Tesseract listo | "
                    f"Versión: {version} | "
                    "Idioma: español"
                )
            )

        except Exception as e:

            self.estado.set(
                "Error verificando Tesseract."
            )

            messagebox.showerror(
                "Error Tesseract",
                str(e)
            )


    # ========================================================
    # SELECCIONAR PDF
    # ========================================================

    def seleccionar_pdf(self):

        archivo = (
            filedialog.askopenfilename(
                title="Seleccionar PDF",
                filetypes=[
                    (
                        "Archivo PDF",
                        "*.pdf"
                    )
                ]
            )
        )

        if not archivo:
            return

        self.pdf_path.set(
            archivo
        )

        nombre = (
            os.path.splitext(
                os.path.basename(
                    archivo
                )
            )[0]
        )

        carpeta = (
            os.path.dirname(
                archivo
            )
        )

        salida = os.path.join(
            carpeta,
            f"{nombre}_OCR.docx"
        )

        self.word_path.set(
            salida
        )


    # ========================================================
    # SELECCIONAR WORD
    # ========================================================

    def seleccionar_word(self):

        archivo = (
            filedialog.asksaveasfilename(
                title="Guardar Word",
                defaultextension=".docx",
                filetypes=[
                    (
                        "Documento Word",
                        "*.docx"
                    )
                ]
            )
        )

        if archivo:

            self.word_path.set(
                archivo
            )


    # ========================================================
    # INICIAR
    # ========================================================

    def iniciar_conversion(self):

        if self.procesando:
            return

        pdf = (
            self.pdf_path
            .get()
            .strip()
        )

        word = (
            self.word_path
            .get()
            .strip()
        )

        if not pdf:

            messagebox.showwarning(
                "PDF",
                "Seleccione un PDF."
            )

            return

        if not os.path.isfile(pdf):

            messagebox.showerror(
                "PDF",
                "El archivo PDF no existe."
            )

            return

        if not word:

            messagebox.showwarning(
                "Word",
                (
                    "Seleccione dónde "
                    "guardar el Word."
                )
            )

            return

        if not TESSERACT_PATH:

            messagebox.showerror(
                "Tesseract",
                (
                    "No se encontró "
                    "Tesseract OCR."
                )
            )

            return

        try:

            idiomas = (
                pytesseract
                .get_languages(
                    config=""
                )
            )

            if OCR_LANGUAGE not in idiomas:

                messagebox.showerror(
                    "Idioma español",
                    (
                        "Tesseract no tiene "
                        "instalado el idioma "
                        "español 'spa'."
                    )
                )

                return

        except Exception as e:

            messagebox.showerror(
                "Tesseract",
                str(e)
            )

            return

        self.cancelar_proceso = False

        self.procesando = True

        self.progreso.set(0)

        self.btn_convertir.config(
            state="disabled"
        )

        self.btn_cancelar.config(
            state="normal"
        )

        hilo = threading.Thread(
            target=self.procesar_pdf,
            daemon=True
        )

        hilo.start()


    # ========================================================
    # PROCESAR PDF
    # ========================================================

    def procesar_pdf(self):

        pdf_path = (
            self.pdf_path.get()
        )

        word_path = (
            self.word_path.get()
        )

        documento_pdf = None

        try:

            self.actualizar_estado(
                "Abriendo PDF..."
            )

            documento_pdf = fitz.open(
                pdf_path
            )

            total_paginas = len(
                documento_pdf
            )

            documento_word = Document()

            configurar_word(
                documento_word
            )

            escala = DPI / 72

            matriz = fitz.Matrix(
                escala,
                escala
            )

            for indice in range(
                total_paginas
            ):

                if self.cancelar_proceso:

                    self.actualizar_estado(
                        "Proceso cancelado."
                    )

                    return

                numero = indice + 1

                self.actualizar_estado(
                    (
                        f"Procesando página "
                        f"{numero} de "
                        f"{total_paginas}..."
                    )
                )

                pagina = (
                    documento_pdf[
                        indice
                    ]
                )

                # =================================================
                # Primero intentamos detectar texto real
                # =================================================

                texto_existente = (
                    pagina.get_text(
                        "text"
                    ).strip()
                )

                # Si la página ya tiene suficiente texto,
                # no hacemos OCR innecesariamente.
                if len(texto_existente) > 100:

                    texto = limpiar_texto(
                        texto_existente
                    )

                else:

                    # Renderizar página
                    pix = (
                        pagina.get_pixmap(
                            matrix=matriz,
                            alpha=False
                        )
                    )

                    # Preparar imagen
                    imagen = (
                        preparar_imagen(
                            pix
                        )
                    )

                    # OCR
                    texto = ejecutar_ocr(
                        imagen
                    )

                    # Liberar memoria
                    del imagen
                    del pix

                # =================================================
                # MARCADOR DE PÁGINA
                # =================================================

                marcador = (
                    documento_word
                    .add_paragraph()
                )

                marcador.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

                run = marcador.add_run(
                    f"— Página {numero} —"
                )

                run.italic = True

                run.font.size = Pt(8)

                # =================================================
                # CONTENIDO
                # =================================================

                if texto:

                    agregar_texto_word(
                        documento_word,
                        texto
                    )

                else:

                    parrafo = (
                        documento_word
                        .add_paragraph()
                    )

                    run = (
                        parrafo.add_run(
                            (
                                "[No se detectó "
                                "texto]"
                            )
                        )
                    )

                    run.italic = True

                # =================================================
                # SALTO DE PÁGINA
                # =================================================

                if numero < total_paginas:

                    documento_word.add_page_break()

                porcentaje = (
                    numero /
                    total_paginas
                ) * 100

                self.actualizar_progreso(
                    porcentaje
                )

            # =====================================================
            # GUARDAR
            # =====================================================

            self.actualizar_estado(
                "Guardando documento Word..."
            )

            carpeta = (
                os.path.dirname(
                    word_path
                )
            )

            if carpeta:

                os.makedirs(
                    carpeta,
                    exist_ok=True
                )

            documento_word.save(
                word_path
            )

            self.actualizar_progreso(
                100
            )

            self.actualizar_estado(
                "Conversión completada."
            )

            self.root.after(
                0,
                lambda: messagebox.showinfo(
                    "Proceso terminado",
                    (
                        "El PDF fue convertido "
                        "correctamente.\n\n"
                        f"Word generado:\n"
                        f"{word_path}"
                    )
                )
            )

        except Exception as e:

            error = str(e)

            self.actualizar_estado(
                "Error durante el proceso."
            )

            self.root.after(
                0,
                lambda: messagebox.showerror(
                    "Error",
                    error
                )
            )

        finally:

            if documento_pdf:

                documento_pdf.close()

            self.procesando = False

            self.root.after(
                0,
                self.restaurar_botones
            )


    # ========================================================
    # ACTUALIZAR UI
    # ========================================================

    def actualizar_estado(
        self,
        texto
    ):

        self.root.after(
            0,
            lambda: self.estado.set(
                texto
            )
        )


    def actualizar_progreso(
        self,
        valor
    ):

        self.root.after(
            0,
            lambda: self.progreso.set(
                valor
            )
        )


    def restaurar_botones(
        self
    ):

        self.btn_convertir.config(
            state="normal"
        )

        self.btn_cancelar.config(
            state="disabled"
        )


    # ========================================================
    # CANCELAR
    # ========================================================

    def cancelar(self):

        if self.procesando:

            self.cancelar_proceso = True

            self.estado.set(
                (
                    "Cancelando después "
                    "de finalizar la "
                    "página actual..."
                )
            )


# ============================================================
# MAIN
# ============================================================

def main():

    root = tk.Tk()

    app = AplicacionOCR(
        root
    )

    root.mainloop()


if __name__ == "__main__":

    main()